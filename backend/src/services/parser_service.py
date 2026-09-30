from __future__ import annotations

import time
import logging
from pathlib import Path
from typing import List

import fitz  # PyMuPDF
import pandas as pd
from docx import Document as DocxDocument

from src.core.logger import logger
from src.domain.schemas import Document

# Supported file extensions for repository ingestion
SUPPORTED_EXTENSIONS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".cpp", ".c", ".h", ".hpp",
    ".cs", ".go", ".rs", ".php", ".rb", ".swift", ".kt", ".md", ".txt", ".json",
    ".yaml", ".yml",
}

LANGUAGE_MAP = {
    ".py": "python",
    ".js": "javascript",
    ".jsx": "javascript",
    ".ts": "typescript",
    ".tsx": "typescript",
    ".java": "java",
    ".cpp": "cpp",
    ".c": "c",
    ".h": "c",
    ".hpp": "cpp",
    ".cs": "csharp",
    ".go": "go",
    ".rs": "rust",
    ".php": "php",
    ".rb": "ruby",
    ".swift": "swift",
    ".kt": "kotlin",
    ".json": "json",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".md": "markdown",
    ".txt": "text",
}

IGNORED_DIRECTORIES = {
    ".git", ".github", ".idea", ".vscode", "__pycache__", "node_modules",
    "venv", ".venv", "dist", "build", "target", ".pytest_cache", ".mypy_cache",
}

IGNORED_FILES = {
    "package-lock.json", "yarn.lock", "pnpm-lock.yaml", "Cargo.lock", "poetry.lock",
}


class ParserService:
    """
    ParserService parses incoming files (PDF, CSV, DOCX, TXT) and repositories
    into transient Document objects representing full files or pages.
    """

    def parse_file(self, file_path: str | Path) -> List[Document]:
        """
        Parses a single file into a list of page/file level Document instances.
        """
        start = time.perf_counter()
        file_path = Path(file_path)
        suffix = file_path.suffix.lower()

        logger.info("Parsing file: %s", file_path.name)

        if suffix == ".pdf":
            documents = self._parse_pdf(file_path)
        elif suffix == ".csv":
            documents = self._parse_csv(file_path)
        elif suffix == ".docx":
            documents = self._parse_docx(file_path)
        elif suffix == ".txt":
            documents = self._parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

        logger.info(
            "Parsing finished in %.2fs (%d files/pages)",
            time.perf_counter() - start,
            len(documents)
        )
        return documents

    def parse_repository(self, repo_path: str | Path) -> List[Document]:
        """
        Scans a local repository directory and parses all supported code/text files
        into whole-file Document instances with basic source metadata.
        """
        start = time.perf_counter()
        repo_path = Path(repo_path)
        documents: List[Document] = []
        repository_name = repo_path.name

        logger.info("Parsing repository: %s", repository_name)

        for file_path in repo_path.rglob("*"):
            if not file_path.is_file():
                continue

            if any(part in IGNORED_DIRECTORIES for part in file_path.parts):
                continue

            if file_path.name in IGNORED_FILES:
                continue

            extension = file_path.suffix.lower()
            if extension not in SUPPORTED_EXTENSIONS:
                continue

            try:
                source = file_path.read_text(encoding="utf-8", errors="ignore")
                if not source.strip():
                    continue

                metadata = {
                    "repository": repository_name,
                    "file": str(file_path.relative_to(repo_path)),
                    "source": str(file_path.relative_to(repo_path)),
                    "extension": extension,
                    "language": LANGUAGE_MAP.get(extension, "unknown"),
                    "type": "github",
                }

                documents.append(
                    Document(page_content=source, metadata=metadata)
                )

            except Exception as e:
                logger.error("Failed to process file %s: %s", file_path, e, exc_info=True)
                raise

        logger.info(
            "Parsed repository %s in %.2fs (%d files)",
            repository_name,
            time.perf_counter() - start,
            len(documents)
        )
        return documents

    # ── Internal parsing methods ──────────────────────────────────────────────

    def _parse_pdf(self, file_path: Path) -> List[Document]:
        documents = []
        with fitz.open(file_path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                documents.append(
                    Document(
                        page_content=page.get_text(),
                        metadata={
                            "source": file_path.name,
                            "page": page_number,
                            "type": "pdf"
                        }
                    )
                )
        return documents

    def _parse_csv(self, file_path: Path) -> List[Document]:
        df = pd.read_csv(file_path)
        return [
            Document(
                page_content=df.to_string(index=False),
                metadata={
                    "source": file_path.name,
                    "type": "csv"
                }
            )
        ]

    def _parse_docx(self, file_path: Path) -> List[Document]:
        doc = DocxDocument(file_path)
        return [
            Document(
                page_content="\n".join(paragraph.text for paragraph in doc.paragraphs),
                metadata={
                    "source": file_path.name,
                    "type": "docx"
                }
            )
        ]

    def _parse_txt(self, file_path: Path) -> List[Document]:
        with open(file_path, "r", encoding="utf-8") as file:
            return [
                Document(
                    page_content=file.read(),
                    metadata={
                        "source": file_path.name,
                        "type": "txt"
                    }
                )
            ]
