from pathlib import Path

import fitz  # PyMuPDF
import pandas as pd
from docx import Document
import time

from app.core.logger import logger
from app.models.document import Document as ArchonDocument


def parse_document(file_path: str | Path):
    start = time.perf_counter()

    file_path = Path(file_path)
    suffix = file_path.suffix.lower()

    logger.info("Parsing %s", file_path.name)

    if suffix == ".pdf":
        documents = _parse_pdf(file_path)

    elif suffix == ".csv":
        documents = _parse_csv(file_path)

    elif suffix == ".docx":
        documents = _parse_docx(file_path)

    elif suffix == ".txt":
        documents = _parse_txt(file_path)

    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    logger.info(
        "Parsing finished in %.2fs",
        time.perf_counter() - start,
    )

    return documents
def _parse_pdf(file_path: Path) -> list[ArchonDocument]:
    documents = []

    with fitz.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            documents.append(
                ArchonDocument(
                    page_content=page.get_text(),
                    metadata={
                        "source": file_path.name,
                        "page": page_number,
                        "type": "pdf"
                    }
                )
            )

    return documents


def _parse_csv(file_path: Path) -> list[ArchonDocument]:
    df = pd.read_csv(file_path)

    return [
        ArchonDocument(
            page_content=df.to_string(index=False),
            metadata={
                "source": file_path.name,
                "type": "csv"
            }
        )
    ]


def _parse_docx(file_path: Path) -> list[ArchonDocument]:
    doc = Document(file_path)

    return [
        ArchonDocument(
            page_content="\n".join(
                paragraph.text
                for paragraph in doc.paragraphs
            ),
            metadata={
                "source": file_path.name,
                "type": "docx"
            }
        )
    ]


def _parse_txt(file_path: Path) -> list[ArchonDocument]:
    with open(file_path, "r", encoding="utf-8") as file:
        return [
            ArchonDocument(
                page_content=file.read(),
                metadata={
                    "source": file_path.name,
                    "type": "txt"
                }
            )
        ]